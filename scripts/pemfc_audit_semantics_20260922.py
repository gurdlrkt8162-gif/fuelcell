#!/usr/bin/env python3
from pathlib import Path
import json,requests,hashlib,zipfile,subprocess,re,io,sys,traceback
from urllib.parse import urljoin
from bs4 import BeautifulSoup
from playwright.sync_api import sync_playwright
from docx import Document
import numpy as np,pandas as pd,scipy.io as sio
R=Path('candidate_bytes');O=Path('pemfc_semantic_audit');O.mkdir(exist_ok=True)
S=requests.Session();S.headers['User-Agent']='Mozilla/5.0 (public academic data verification)'
LOG=[]

def write(name,obj):
    (O/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2,allow_nan=False,default=str),encoding='utf-8')

def fetch(url,p,expected=None,md5=None):
    p.parent.mkdir(parents=True,exist_ok=True);r=S.get(url,timeout=(20,90));r.raise_for_status();b=r.content
    if len(b)>100000000 or b.lstrip().lower().startswith((b'<!doctype html',b'<html')):raise ValueError('oversize or HTML')
    if expected is not None and len(b)!=expected:raise ValueError(f'byte mismatch {len(b)} != {expected}')
    if md5 and hashlib.md5(b).hexdigest()!=md5:raise ValueError('MD5 mismatch')
    p.write_bytes(b);LOG.append({'url':url,'path':str(p),'size':len(b),'sha256':hashlib.sha256(b).hexdigest(),'provider_md5_verified':bool(md5)})
    return b

def unpack(p):
    q=p.with_name(p.name+'_unpacked');q.mkdir(exist_ok=True)
    t=subprocess.run(['7z','t',str(p)],capture_output=True,text=True,timeout=90)
    if t.returncode:raise ValueError(t.stdout[-1000:])
    x=subprocess.run(['7z','x','-y','-o'+str(q),str(p)],capture_output=True,text=True,timeout=90)
    if x.returncode:raise ValueError(x.stdout[-1000:])
    return q

with sync_playwright() as w:
    browser=w.chromium.launch(headless=True)
    for tag,id_,ver in [('RWTH','mc46tw9t8m','1'),('OXYGEN_PURGE','jy5vxystwd','2')]:
        ctx=browser.new_context(accept_downloads=True);page=ctx.new_page();net=[]
        def response(resp):
            if 'public-api' in resp.url and resp.status==200 and 'json' in resp.headers.get('content-type',''):
                try:net.append({'url':resp.url,'body':resp.json()})
                except Exception:pass
        page.on('response',response)
        try:
            u=f'https://data.mendeley.com/datasets/{id_}/{ver}';page.goto(u,wait_until='networkidle',timeout=120000)
            for label in ['Accept All Cookies','Accept all cookies','I agree']:
                try:page.get_by_role('button',name=label,exact=False).first.click(timeout=800);break
                except Exception:pass
            with page.expect_download(timeout=180000) as info:page.get_by_text('Download All',exact=False).first.click()
            d=info.value;p=R/tag/(id_+'.zip');p.parent.mkdir(parents=True,exist_ok=True);d.save_as(str(p))
            if d.failure():raise ValueError(d.failure())
            b=p.read_bytes();LOG.append({'tag':tag,'source':u,'file':str(p),'size':len(b),'sha256':hashlib.sha256(b).hexdigest(),'source_download_url':d.url,'state':'BYTES'})
            unpack(p)
        except Exception as e:LOG.append({'tag':tag,'error':str(e)})
        write(tag+'_browser_download.json',{'events':net,'page_text':page.locator('body').inner_text()[:30000]})
        ctx.close()
    # Publisher supplementary link discovery in a normal public browser.
    ctx=browser.new_context(accept_downloads=True);page=ctx.new_page()
    try:
        u='https://www.mdpi.com/2306-5729/5/2/47';page.goto(u,wait_until='domcontentloaded',timeout=60000);page.wait_for_timeout(3000)
        links=page.locator('a').evaluate_all('(es)=>es.map(e=>({text:e.innerText,url:e.href}))');write('SEVILLA_links.json',links)
        candidates=[x for x in links if '.rar' in x['url'].lower() or '.zip' in x['url'].lower()]
        if not candidates:
            for x in links:
                if x['url'].endswith('/s1'):
                    page.goto(x['url'],wait_until='domcontentloaded',timeout=60000)
                    candidates=page.locator('a').evaluate_all('(es)=>es.map(e=>({text:e.innerText,url:e.href}))');candidates=[x for x in candidates if '.rar' in x['url'].lower() or '.zip' in x['url'].lower()];break
        for i,x in enumerate(candidates):
            p=R/'SEVILLA'/('original_'+str(i)+('.rar' if '.rar' in x['url'].lower() else '.zip'));fetch(x['url'],p);unpack(p)
        if not candidates:LOG.append({'tag':'SEVILLA','error':'No accessible supplementary archive link found','page_text':page.locator('body').inner_text()[:2000]})
    except Exception as e:LOG.append({'tag':'SEVILLA','error':str(e)})
    ctx.close();browser.close()
# Dataverse's tabular ingest changes stored size/checksum. Match the format actually requested.
meta=json.loads(Path('pemfc_candidate_audit/DIFFERENTIAL_source.json').read_text())
filelist=meta['data']['latestVersion']['files'];write('DIFFERENTIAL_file_metadata.json',filelist)
for f in filelist:
    d=f['dataFile'];p=R/'DIFFERENTIAL'/f.get('directoryLabel','')/d['filename']
    if p.exists():continue
    u='https://entrepot.recherche.data.gouv.fr/api/access/datafile/'+str(d['id'])
    try:
        # Default is the archived/tabular representation, matching file metadata.
        fetch(u,p,d['filesize'],d.get('md5'))
    except Exception as e:LOG.append({'tag':'DIFFERENTIAL','file':d['filename'],'error':str(e)})
# Expand archive payload only, never execute bundled author scripts.
for level in range(3):
    for p in list(R.rglob('*')):
        if p.is_file() and p.suffix.lower() in ('.7z','.rar','.zip') and not p.with_name(p.name+'_unpacked').exists():
            try:unpack(p)
            except Exception as e:LOG.append({'file':str(p),'error':str(e)})
# Snapshot channel dictionaries and human-readable methods.
for tag in [p for p in R.iterdir() if p.is_dir()]:
    for p in tag.rglob('*'):
        if not p.is_file():continue
        dest=re.sub('[^A-Za-z0-9._-]','_',str(p.relative_to(tag)))
        try:
            if p.suffix.lower()=='.docx':
                d=Document(p);text='\n'.join(x.text for x in d.paragraphs)+'\n'+ '\n\n'.join('\n'.join('\t'.join(c.text for c in r.cells) for r in t.rows) for t in d.tables)
                (O/(tag.name+'__'+dest+'.txt')).write_text(text,encoding='utf-8')
            elif p.suffix.lower() in ('.txt','.m','.md','.tab') and p.stat().st_size<65000:
                if p.suffix.lower()!='.txt' or any(k in p.name.lower() for k in ('readme','info','condition')) or tag.name=='RWTH':
                    (O/(tag.name+'__'+dest)).write_text(p.read_bytes().decode('utf-8-sig',errors='replace'),encoding='utf-8')
        except Exception as e:LOG.append({'path':str(p),'text_error':str(e)})
# Numeric content summary, not acceptance by workflow exit code.
for tag in [p for p in R.iterdir() if p.is_dir()]:
    rows=[]
    for p in tag.rglob('*'):
        if not p.is_file():continue
        base={'path':str(p),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
        try:
            if p.suffix.lower() in ('.csv','.txt','.tab','.dat','.asc'):
                txt=p.read_bytes().decode('utf-8-sig',errors='replace');base['first_lines']=txt.splitlines()[:10]
                try:
                    frame=pd.read_csv(io.StringIO(txt),sep=None,engine='python')
                    base.update({'rows':len(frame),'cols':len(frame.columns),'headers':[str(x) for x in frame.columns]})
                except Exception as e:base['parse_error']=str(e)
            elif p.suffix.lower() in ('.xls','.xlsx'):
                xf=pd.ExcelFile(p);sheets=[]
                for n in xf.sheet_names:
                    df=pd.read_excel(p,sheet_name=n,header=None);a=df.apply(pd.to_numeric,errors='coerce').to_numpy();valid=np.isfinite(a)
                    sheets.append({'sheet':n,'rows':len(df),'cols':len(df.columns),'missing':int(df.isna().sum().sum()),'numeric':int(valid.sum()),'head':df.head(3).fillna('').astype(str).values.tolist(),'min_by_col':[float(a[:,i][valid[:,i]].min()) if valid[:,i].any() else None for i in range(a.shape[1])],'max_by_col':[float(a[:,i][valid[:,i]].max()) if valid[:,i].any() else None for i in range(a.shape[1])]})
                base['sheets']=sheets
            elif p.suffix.lower()=='.mat':
                base['mat_variables']=sio.whosmat(p)
                d=sio.loadmat(p,squeeze_me=True,struct_as_record=False)
                keys=[]
                def visit(x,name,depth=0):
                    if depth>6:return
                    if isinstance(x,dict):
                        for k,v in x.items():
                            if not k.startswith('__'):visit(v,name+'.'+k,depth+1)
                    elif hasattr(x,'_fieldnames'):
                        for k in x._fieldnames:visit(getattr(x,k),name+'.'+k,depth+1)
                    elif isinstance(x,np.ndarray):
                        r={'key':name,'shape':list(x.shape),'dtype':str(x.dtype)}
                        if x.dtype.kind in 'biuf':
                            a=x.astype(float);ok=np.isfinite(a);r.update({'nan':int(np.isnan(a).sum()),'inf':int(np.isinf(a).sum()),'min':float(a[ok].min()) if ok.any() else None,'max':float(a[ok].max()) if ok.any() else None,'head':a.ravel()[:5].tolist()})
                            if a.ndim==1 and ('time' in name.lower()):
                                dt=np.diff(a);r['time_differences']={'negative':int((dt<0).sum()),'duplicates':int((dt==0).sum()),'median':float(np.median(dt)),'min':float(dt.min()),'max':float(dt.max())}
                        elif x.dtype.kind in 'US':r['head']=x.ravel()[:50].tolist()
                        keys.append(r)
                    else:keys.append({'key':name,'value':str(x)[:200]})
                visit(d,'mat');base['mat_fields']=keys
            else:continue
        except Exception as e:base['error']=str(e)
        rows.append(base)
    write(tag.name+'_OVERVIEW.json',rows)
write('DOWNLOAD_AND_PARSE_LOG.json',LOG)
with zipfile.ZipFile('PEMFC_STAGE3_CANDIDATE_NOT_APPROVED.zip','w',zipfile.ZIP_DEFLATED,compresslevel=3) as z:
    for p in R.rglob('*'):
        if p.is_file() and '_unpacked' not in str(p) and not p.name.endswith('.part'):z.write(p,p)
    for p in O.rglob('*'):
        if p.is_file():z.write(p,p)
print('SEMANTIC_INSPECTION_COMPLETE',json.dumps(LOG,ensure_ascii=False)[:12000])
