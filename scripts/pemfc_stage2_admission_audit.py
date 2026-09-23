#!/usr/bin/env python3
"""Strict second-stage audit of public PEMFC datasets.

Downloads only public files, verifies provider size/checksum when exposed,
tests archives, parses numerical schemas, and records quality defects.  No
candidate is automatically admitted to the user's research archive.
"""
from __future__ import annotations

from pathlib import Path
from urllib.parse import urljoin
import hashlib, io, json, math, re, subprocess, time, zipfile

import h5py
import numpy as np
import pandas as pd
import requests
import scipy.io as sio
from bs4 import BeautifulSoup
from docx import Document

RAW = Path("pemfc_stage2_raw")
AUD = Path("pemfc_stage2_audit")
RAW.mkdir(exist_ok=True); AUD.mkdir(exist_ok=True)
S = requests.Session()
S.headers.update({"User-Agent":"PEMFC-public-data-admission-audit/2026-09-23", "Accept":"*/*"})
RECEIPTS=[]; ERRORS=[]


def dump(name, obj):
    p=AUD/name; p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str),encoding="utf-8")


def hash_file(path, alg="sha256"):
    h=hashlib.new(alg)
    with Path(path).open("rb") as f:
        for b in iter(lambda:f.read(8*1024*1024),b""): h.update(b)
    return h.hexdigest()


def get(url, stream=False):
    last=None
    for k in range(6):
        try:
            r=S.get(url,stream=stream,timeout=(30,240),allow_redirects=True)
            if r.status_code in (429,500,502,503,504):
                last=r; time.sleep(min(30,2**k)); continue
            return r
        except requests.RequestException as e:
            last=e; time.sleep(min(30,2**k))
    if isinstance(last,requests.Response): return last
    raise last


def safe(s):
    return (re.sub(r"[^A-Za-z0-9._()\-+ ]+","_",str(s or "download")).strip(" .")[:220] or "download")


def download(url,dest,size=None,sha256=None,md5=None,strict_size=True):
    dest=Path(dest); dest.parent.mkdir(parents=True,exist_ok=True)
    tmp=dest.with_name(dest.name+".part"); tmp.unlink(missing_ok=True)
    r=get(url,stream=True)
    if r.status_code!=200: raise RuntimeError(f"HTTP {r.status_code} {url}")
    ct=(r.headers.get("content-type") or "").lower(); n=0; h2=hashlib.sha256(); hm=hashlib.md5(); head=b""
    with tmp.open("wb") as f:
        for b in r.iter_content(2*1024*1024):
            if not b: continue
            if len(head)<1024: head+=b[:1024-len(head)]
            n+=len(b)
            if n>400*1024*1024: raise RuntimeError("per-file 400 MiB audit gate")
            f.write(b); h2.update(b); hm.update(b)
    if not n or head.lstrip().lower().startswith((b"<html",b"<!doctype html")) or "text/html" in ct:
        raise RuntimeError("empty/HTML response")
    if size not in (None,"",0,"0") and strict_size and n!=int(size): raise RuntimeError(f"size {n}!={size}")
    g2,gm=h2.hexdigest(),hm.hexdigest()
    if sha256 and g2.lower()!=sha256.lower(): raise RuntimeError("sha256 mismatch")
    if md5 and gm.lower()!=md5.lower(): raise RuntimeError("md5 mismatch")
    tmp.replace(dest)
    row={"path":str(dest),"url":url,"final_url":r.url,"bytes":n,"content_type":ct,"sha256":g2,"md5":gm,
         "expected_size":size,"expected_sha256":sha256,"expected_md5":md5,
         "size_match":None if size in (None,"",0,"0") else n==int(size),
         "sha256_match":None if not sha256 else g2.lower()==sha256.lower(),
         "md5_match":None if not md5 else gm.lower()==md5.lower()}
    RECEIPTS.append(row); return row


def sevenzip(path):
    path=Path(path); out=path.with_name(path.name+"__unpacked")
    if out.exists(): return out
    out.mkdir(parents=True)
    t=subprocess.run(["7z","t",str(path)],capture_output=True,text=True,timeout=240)
    if t.returncode: raise RuntimeError("archive CRC/integrity failure: "+t.stdout[-1500:]+t.stderr[-500:])
    ls=subprocess.run(["7z","l","-slt",str(path)],capture_output=True,text=True,timeout=120).stdout
    for line in ls.splitlines():
        if line.startswith("Path = "):
            x=line[7:].replace("\\","/")
            if x.startswith("/") or ".." in Path(x).parts: raise RuntimeError("unsafe archive member "+x)
    x=subprocess.run(["7z","x","-y",f"-o{out}",str(path)],capture_output=True,text=True,timeout=300)
    if x.returncode: raise RuntimeError("archive extraction failure: "+x.stdout[-1500:]+x.stderr[-500:])
    return out


def mendeley_listing(ds,ver):
    base="https://data.mendeley.com/public-api"
    vr=get(f"{base}/datasets/{ds}/versions"); vr.raise_for_status()
    fr=get(f"{base}/datasets/{ds}/folders/{ver}"); fr.raise_for_status(); folders=fr.json()
    todo=[("root","root")]+[(f.get("name") or f["id"],f["id"]) for f in folders]
    rows=[]
    for name,fid in todo:
        r=get(f"{base}/datasets/{ds}/files?folder_id={fid}&version={ver}"); r.raise_for_status()
        for q in r.json(): q=dict(q); q["folder_name"]=name; rows.append(q)
    return {"versions":vr.json(),"folders":folders,"files":rows}


def take_mendeley(tag,ds,ver):
    meta=mendeley_listing(ds,ver); dump(f"{tag}_mendeley_listing.json",meta)
    if not meta["files"]: raise RuntimeError("no files resolved")
    for q in meta["files"]:
        c=q.get("content_details") or {}; u=c.get("download_url")
        if not u: raise RuntimeError("missing file URL "+str(q.get("filename")))
        download(u,RAW/tag/safe(q.get("folder_name"))/safe(q.get("filename")),size=c.get("size") or q.get("size"),sha256=c.get("sha256_hash"))


def take_figshare(tag,article):
    r=get(f"https://api.figshare.com/v2/articles/{article}"); r.raise_for_status(); meta=r.json()
    dump(f"{tag}_figshare_metadata.json",meta)
    for q in meta.get("files",[]): download(q["download_url"],RAW/tag/safe(q["name"]),size=q.get("size"),md5=q.get("supplied_md5") or q.get("computed_md5"))


def take_differential():
    api="https://entrepot.recherche.data.gouv.fr/api/datasets/:persistentId/?persistentId=doi:10.57745/YLCIPH"
    r=get(api); r.raise_for_status(); meta=r.json(); dump("DIFFERENTIAL_metadata.json",meta)
    for q in meta["data"]["latestVersion"]["files"]:
        if q.get("restricted"): raise RuntimeError("restricted file")
        d=q["dataFile"]; fid=d["id"]; label=d.get("filename") or q.get("label")
        dest=RAW/"DIFFERENTIAL"/(q.get("directoryLabel") or "root")/safe(label)
        u=f"https://entrepot.recherche.data.gouv.fr/api/access/datafile/{fid}"
        try: download(u,dest,size=d.get("filesize"),md5=(d.get("checksum") or {}).get("value"))
        except Exception as e:
            orig=dest.with_name(dest.stem+"__original"+dest.suffix)
            rec=download(u+"?format=original",orig,strict_size=False)
            rec["dataverse_representation_issue"]=str(e)
            ERRORS.append({"tag":"DIFFERENTIAL","file":label,"nonfatal":True,"issue":"Dataverse representation mismatch","detail":str(e),"preserved":str(orig)})


def take_sevilla():
    dst=RAW/"SEVILLA_DATA2020"/"Experimental_performance_load_cycling_PEMFC.rar"
    official="https://www.mdpi.com/2306-5729/5/2/47/s1"
    try: download(official,dst); return
    except Exception as first:
        candidates=["https://mdpi-res.com/d_attachment/data/data-05-00047/article_deploy/data-05-00047-s001.rar"]
        page=get("https://www.mdpi.com/2306-5729/5/2/47")
        if page.status_code==200 and "html" in (page.headers.get("content-type") or ""):
            soup=BeautifulSoup(page.text,"html.parser")
            candidates=[urljoin(page.url,a.get("href")) for a in soup.find_all("a",href=True) if "s001" in a.get("href","").lower() or a.get("href","").lower().endswith(".rar")]+candidates
        last=first
        for u in dict.fromkeys(candidates):
            try:
                download(u,dst); dump("SEVILLA_resolution.json",{"official":official,"resolved":u,"initial_error":str(first)}); return
            except Exception as e: last=e
        raise RuntimeError("official supplement unresolved: "+str(last))


def finite(v):
    try: x=float(v); return x if math.isfinite(x) else None
    except Exception: return None


def profile(df,header="parsed"):
    out={"rows":len(df),"columns":len(df.columns),"headers":[str(x) for x in df.columns],"header_mode":header,"duplicate_rows":int(df.duplicated().sum()),"stats":[]}
    for i,c in enumerate(df.columns):
        s=df.iloc[:,i]; n=pd.to_numeric(s,errors="coerce").to_numpy(dtype=float); ok=np.isfinite(n); v=n[ok]
        q={"index":i,"name":str(c),"numeric":int(ok.sum()),"missing":int(s.isna().sum()),"unique":int(s.nunique(dropna=True)),
           "min":finite(v.min()) if len(v) else None,"max":finite(v.max()) if len(v) else None,"mean":finite(v.mean()) if len(v) else None,"examples":[str(x) for x in s.head(4)]}
        if i==0 and len(v)>2:
            d=np.diff(v); q["diff"]={"negative":int((d<0).sum()),"zero":int((d==0).sum()),"median":finite(np.median(d)),"min":finite(d.min()),"max":finite(d.max())}
        out["stats"].append(q)
    return out


def mat_desc(x,depth=0):
    if depth>5:return {"type":type(x).__name__}
    if isinstance(x,dict):return {k:mat_desc(v,depth+1) for k,v in x.items() if not k.startswith("__")}
    if hasattr(x,"_fieldnames"):return {k:mat_desc(getattr(x,k),depth+1) for k in x._fieldnames}
    if isinstance(x,np.ndarray):
        r={"shape":list(x.shape),"dtype":str(x.dtype)}
        if x.dtype.kind in "biuf" and x.size:
            a=x.astype(float); ok=np.isfinite(a); r.update({"finite":int(ok.sum()),"nonfinite":int((~ok).sum()),"min":finite(a[ok].min()) if ok.any() else None,"max":finite(a[ok].max()) if ok.any() else None})
        elif x.dtype.kind=="O":r["first"]=[mat_desc(v,depth+1) for v in x.ravel()[:4]]
        return r
    if isinstance(x,(int,float,np.number)):return finite(x)
    return str(x)[:500]


def inspect(path):
    path=Path(path); r={"path":str(path),"bytes":path.stat().st_size,"sha256":hash_file(path)}; ext=path.suffix.lower()
    try:
        if ext in (".zip",".rar",".7z"):
            root=sevenzip(path); r["archive_test"]="PASS"; r["members"]=[{"path":str(p.relative_to(root)),"bytes":p.stat().st_size} for p in root.rglob("*") if p.is_file()]
        elif ext in (".xlsx",".xls"):
            x=pd.ExcelFile(path); r["sheets"]={}
            for n in x.sheet_names:
                a=pd.read_excel(path,sheet_name=n); b=pd.read_excel(path,sheet_name=n,header=None)
                r["sheets"][n]={"inferred":profile(a),"header_none":profile(b,"none")}
        elif ext==".docx":
            d=Document(path); r["paragraphs"]=[p.text for p in d.paragraphs if p.text.strip()]; r["tables"]=[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
        elif ext==".mat":
            try:r["mat"]=mat_desc(sio.loadmat(path,squeeze_me=True,struct_as_record=False))
            except NotImplementedError:
                with h5py.File(path) as f:
                    z=[]; f.visititems(lambda n,o:z.append({"name":n,"shape":list(o.shape),"dtype":str(o.dtype)}) if isinstance(o,h5py.Dataset) else None); r["hdf5"]=z
        elif ext in (".csv",".txt",".tab",".dat",".asc",".m",".md",".names"):
            text=path.read_bytes().decode("utf-8-sig",errors="replace"); r["line_count"]=len(text.splitlines()); r["preview"]=text[:18000]
            if ext not in (".m",".md",".names") and r["line_count"]>4:
                for sep in (None,"\t",";",","):
                    try:
                        df=pd.read_csv(path,sep=sep,engine="python")
                        if len(df.columns)>1:r["table"]=profile(df); break
                    except Exception:pass
        else:r["state"]="binary-unparsed"
    except Exception as e:r["error"]=repr(e)
    return r


def expand_archives():
    for _ in range(5):
        todo=[p for p in RAW.rglob("*") if p.is_file() and p.suffix.lower() in (".zip",".rar",".7z") and not p.with_name(p.name+"__unpacked").exists()]
        if not todo:break
        for p in todo:
            try:sevenzip(p)
            except Exception as e:ERRORS.append({"path":str(p),"archive_error":repr(e)})


def special_lough2018():
    fs=list((RAW/"LOUGH2018").rglob("*.CSV"))+list((RAW/"LOUGH2018").rglob("*.csv"))
    if len(fs)!=1:raise RuntimeError(f"expected one CSV got {len(fs)}")
    df=pd.read_csv(fs[0]); out={"profile":profile(df),"range_flags":{},"consistency":{}}
    ranges={"U_tot (V)":(0,2),"RH_Air":(0,100),"RH_H2":(0,100),"T_1":(-20,150),"T_2":(-20,150),"T_3":(-20,150),"T_4":(-20,150),"T_Air_inlet":(-20,150),"T_H2_inlet":(-20,150),"T_Stack_outlet":(-20,150),"T_Heater":(-20,500)}
    for c,(lo,hi) in ranges.items():
        if c in df:
            x=pd.to_numeric(df[c],errors="coerce"); out["range_flags"][c]={"below":int((x<lo).sum()),"above":int((x>hi).sum()),"nonfinite":int((~np.isfinite(x)).sum())}
    if {"P (W)","U_tot (V)","i (A)"}.issubset(df.columns):
        v=df["U_tot (V)"].between(0,2)&np.isfinite(df[["P (W)","U_tot (V)","i (A)"]]).all(axis=1); e=df.loc[v,"P (W)"]-df.loc[v,"U_tot (V)"]*df.loc[v,"i (A)"]
        out["consistency"]["P_minus_UI"]={"n":len(e),"rmse_W":finite(np.sqrt(np.mean(e**2))),"median_abs_W":finite(np.median(np.abs(e))),"p99_abs_W":finite(np.quantile(np.abs(e),.99))}
    if {"i (A)","i_write"}.issubset(df.columns):
        e=df["i (A)"]-df["i_write"]; out["consistency"]["current_tracking"]={"rmse_A":finite(np.sqrt(np.nanmean(e**2))),"median_abs_A":finite(np.nanmedian(np.abs(e))),"p99_abs_A":finite(np.nanquantile(np.abs(e),.99))}
    dump("LOUGH2018_QC.json",out)


def main():
    tasks=[
      ("RWTH",lambda:take_mendeley("RWTH","mc46tw9t8m",1)),
      ("OXYGEN_PURGE",lambda:take_mendeley("OXYGEN_PURGE","jy5vxystwd",2)),
      ("TWO_SYSTEM_FAULT",lambda:take_mendeley("TWO_SYSTEM_FAULT","m9m2gmcr6k",1)),
      ("SENSOR_FAULT",lambda:take_mendeley("SENSOR_FAULT","4sfyzcc39j",1)),
      ("LOUGH2016",lambda:take_figshare("LOUGH2016",4009959)),
      ("LOUGH2018",lambda:take_figshare("LOUGH2018",5759667)),
      ("DIFFERENTIAL",take_differential),("SEVILLA_DATA2020",take_sevilla)]
    for tag,fn in tasks:
        print("==",tag,"==",flush=True)
        try:fn()
        except Exception as e:ERRORS.append({"tag":tag,"task_error":repr(e)}); print("ERROR",tag,e,flush=True)
    expand_archives()
    for root in sorted(p for p in RAW.iterdir() if p.is_dir()):
        profiles=[]
        for p in sorted(root.rglob("*")):
            if not p.is_file() or p.name.endswith(".part"):continue
            profiles.append({"path":str(p),"bytes":p.stat().st_size,"sha256":hash_file(p),"state":"parser-size-gate"} if p.stat().st_size>180*1024*1024 else inspect(p))
        dump(f"{root.name}_PROFILES.json",profiles)
        dump(f"{root.name}_SUMMARY.json",{"files":len(profiles),"parse_errors":[x for x in profiles if "error" in x],"bytes_including_extracted":sum(int(x.get("bytes",0)) for x in profiles)})
    try:special_lough2018()
    except Exception as e:ERRORS.append({"tag":"LOUGH2018","qc_error":repr(e)})
    dump("DOWNLOAD_RECEIPTS.json",RECEIPTS); dump("ERRORS.json",ERRORS)
    dump("AUDIT_SUMMARY.json",{"datasets_requested":[x[0] for x in tasks],"downloaded_files":len(RECEIPTS),"downloaded_bytes":sum(x["bytes"] for x in RECEIPTS),"errors":ERRORS,"state":"PARSED_FOR_MANUAL_ADMISSION_REVIEW_NOT_APPROVED"})
    out=Path("PEMFC_STAGE2_CANDIDATES_NOT_APPROVED.zip")
    with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED,compresslevel=4,allowZip64=True) as z:
        for p in RAW.rglob("*"):
            if p.is_file() and "__unpacked" not in str(p) and not p.name.endswith(".part"):z.write(p,p)
        for p in AUD.rglob("*"):
            if p.is_file():z.write(p,p)
    print(json.dumps({"artifact":str(out),"bytes":out.stat().st_size,"sha256":hash_file(out),"errors":len(ERRORS)},indent=2),flush=True)

if __name__=="__main__":main()
